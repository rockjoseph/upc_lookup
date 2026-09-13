"""Reading, updating, appending to, and exporting the price-tracking
.xlsx workbook via openpyxl.

Each browser session gets its own workbook file on disk under
`data/sessions/<session_id>.xlsx`, created either from an uploaded file or
from scratch. Rows are matched first by SKU/UPC, then falling back to an
exact (case-insensitive) product name match, per the spec.
"""
import os
import uuid
from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import EXCEL_HEADERS, EXCEL_SHEET_NAME, SESSIONS_DIR
from app.models import ExcelRow, Product

COL_SKU = 1
COL_NAME = 2
COL_ORIGINAL_PRICE = 3
COL_SALE_PRICE = 4
COL_PROMO = 5
COL_LAST_CHECKED = 6
COL_URL = 7


class ExcelServiceError(Exception):
    pass


def _session_path(session_id: str) -> str:
    safe_id = "".join(c for c in session_id if c.isalnum() or c in "-_")
    if not safe_id:
        raise ExcelServiceError("Invalid session id")
    return os.path.join(SESSIONS_DIR, f"{safe_id}.xlsx")


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = EXCEL_SHEET_NAME
    ws.append(EXCEL_HEADERS)
    for col_idx, header in enumerate(EXCEL_HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            14, len(header) + 2
        )
    return wb


def _get_sheet(wb: Workbook) -> Worksheet:
    if EXCEL_SHEET_NAME in wb.sheetnames:
        return wb[EXCEL_SHEET_NAME]
    # Uploaded workbook may use a different sheet name / layout; use the
    # first (active) sheet and trust it follows the expected column order.
    return wb.active


def create_session(upload_path: Optional[str] = None) -> str:
    """Creates a new session workbook, optionally seeded from an uploaded
    file. Returns the new session id."""
    session_id = uuid.uuid4().hex
    dest_path = _session_path(session_id)

    if upload_path:
        wb = load_workbook(upload_path)
        ws = _get_sheet(wb)
        header_row = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
        if not header_row or header_row[0] is None:
            # Empty/unlabeled sheet -- stamp our headers onto it.
            ws.delete_rows(1, ws.max_row) if ws.max_row else None
            ws.append(EXCEL_HEADERS)
        wb.save(dest_path)
    else:
        wb = _new_workbook()
        wb.save(dest_path)

    return session_id


def _row_to_excel_row(row) -> Optional[ExcelRow]:
    sku, name = row[COL_SKU - 1].value, row[COL_NAME - 1].value
    if sku is None and name is None:
        return None
    orig = row[COL_ORIGINAL_PRICE - 1].value if len(row) >= COL_ORIGINAL_PRICE else None
    sale = row[COL_SALE_PRICE - 1].value if len(row) >= COL_SALE_PRICE else None
    promo = row[COL_PROMO - 1].value if len(row) >= COL_PROMO else None
    checked = row[COL_LAST_CHECKED - 1].value if len(row) >= COL_LAST_CHECKED else None
    url = row[COL_URL - 1].value if len(row) >= COL_URL else None
    return ExcelRow(
        sku=str(sku) if sku is not None else None,
        name=str(name) if name is not None else "",
        original_price=float(orig) if isinstance(orig, (int, float)) else None,
        sale_price=float(sale) if isinstance(sale, (int, float)) else None,
        promo_deal=str(promo) if promo else None,
        last_checked=str(checked) if checked else None,
        url=str(url) if url else None,
    )


def get_rows(session_id: str) -> List[ExcelRow]:
    path = _session_path(session_id)
    if not os.path.exists(path):
        raise ExcelServiceError(f"Unknown session: {session_id}")
    wb = load_workbook(path)
    ws = _get_sheet(wb)
    rows = []
    for excel_row in ws.iter_rows(min_row=2):
        parsed = _row_to_excel_row(excel_row)
        if parsed:
            rows.append(parsed)
    return rows


def upsert_product(session_id: str, product: Product) -> List[ExcelRow]:
    """Find a row matching the product by SKU (preferred) or name and
    update it; otherwise append a new row. Returns the full updated row
    list so the frontend can refresh its preview table."""
    path = _session_path(session_id)
    if not os.path.exists(path):
        raise ExcelServiceError(f"Unknown session: {session_id}")

    wb = load_workbook(path)
    ws = _get_sheet(wb)

    match_row_idx: Optional[int] = None
    for row in ws.iter_rows(min_row=2):
        sku_cell = row[COL_SKU - 1]
        name_cell = row[COL_NAME - 1]
        if product.sku and sku_cell.value is not None and str(sku_cell.value).strip() == str(product.sku).strip():
            match_row_idx = row[0].row
            break
        if (
            not match_row_idx
            and name_cell.value
            and str(name_cell.value).strip().lower() == product.name.strip().lower()
        ):
            match_row_idx = row[0].row

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if match_row_idx is None:
        match_row_idx = ws.max_row + 1 if ws.max_row >= 1 else 2

    ws.cell(row=match_row_idx, column=COL_SKU, value=product.sku or "")
    ws.cell(row=match_row_idx, column=COL_NAME, value=product.name)
    ws.cell(row=match_row_idx, column=COL_ORIGINAL_PRICE, value=product.original_price)
    ws.cell(row=match_row_idx, column=COL_SALE_PRICE, value=product.sale_price)
    ws.cell(row=match_row_idx, column=COL_PROMO, value=product.promo_deal or "")
    ws.cell(row=match_row_idx, column=COL_LAST_CHECKED, value=timestamp)
    ws.cell(row=match_row_idx, column=COL_URL, value=product.url)

    wb.save(path)
    return get_rows(session_id)


def get_file_path(session_id: str) -> str:
    path = _session_path(session_id)
    if not os.path.exists(path):
        raise ExcelServiceError(f"Unknown session: {session_id}")
    return path


def cleanup_expired_sessions(ttl_seconds: int) -> int:
    """Deletes session workbooks older than ttl_seconds. Returns count removed."""
    import time

    removed = 0
    now = time.time()
    for filename in os.listdir(SESSIONS_DIR):
        path = os.path.join(SESSIONS_DIR, filename)
        try:
            if now - os.path.getmtime(path) > ttl_seconds:
                os.remove(path)
                removed += 1
        except OSError:
            continue
    return removed
